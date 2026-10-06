"""Excel (.xlsx) version of the Customer Version List export — same rows,
same column choices and same exclusions (no Renewal, no Health) as
customer_export_pdf.py, whose COLUMN_ORDER/COLUMN_LABELS this reuses so
the two formats can never drift apart.

Built for working with the list, not just reading it: real typed cells
(numbers stay numbers so they sort and sum), a frozen header with
autofilter, and the PDF's red/amber/green signals carried over as font
colours rather than baked into text. "Upgrades" is split into Used and
Limit columns because "3/4" as text can't be filtered or summed.
"""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.services.customer_export_pdf import COLUMN_ORDER, COLUMN_LABELS, _is_jvm_version

_RED, _AMBER, _GREEN, _MUTED = "C62828", "B26A00", "1B7F4B", "8A94A6"
_HEADER_FILL = PatternFill("solid", fgColor="1C2A45")
_HEADER_FONT = Font(bold=True, color="FFFFFF")

_WIDTHS = {
    "name": 34, "tier": 11, "csm": 18, "package": 16, "prod_version": 13, "infra": 9,
    "days_since_upgrade": 12, "upgrades_used": 10, "upgrades_limit": 10, "open_cases": 10, "migration": 14,
    "primary_contact": 38,
}


def _version_tuple(v: str) -> tuple[int, ...] | None:
    try:
        return tuple(int(p) for p in v.replace("-R", "").split("."))
    except ValueError:
        return None


def _cell_value(key: str, row: dict):
    """Returns (value, font colour or None, number format or None)."""
    if key == "prod_version":
        version = row.get("prod_version")
        if row.get("jvm_client") or _is_jvm_version(version):
            return "JVM", None, None
        if not version:
            return None, None, None
        vt = _version_tuple(version)
        color = _RED if vt and vt < (8, 23) else _AMBER if vt and vt < (8, 27) else None
        return version, color, "@"
    if key == "days_since_upgrade":
        days = row.get("days_since_upgrade")
        if days is None:
            return "Never", _MUTED, None
        return days, _RED if days > 365 else _AMBER if days > 90 else _GREEN, "0"
    if key == "upgrades_used":
        used, limit = row.get("upgrades_used") or 0, row.get("upgrades_limit") or 0
        return used, _RED if limit and used >= limit else _AMBER if limit and used > limit * 0.7 else None, "0"
    if key == "upgrades_limit":
        return row.get("upgrades_limit") or 0, None, "0"
    if key == "open_cases":
        return row.get("open_cases"), None, "0"
    if key == "migration":
        stage = row.get("migration_stage")
        return stage, _GREEN if stage == "Complete" else None, None
    if key == "primary_contact":
        # Joined rather than first-only: a few customers have more than one
        # contact flagged primary, and silently dropping one would hide that.
        return "; ".join(row.get("primary_contacts") or []) or None, None, None
    return row.get(key) or None, None, None


def render_xlsx(customers: list[dict], columns: list[str], generated_at: str) -> bytes:
    cols: list[str] = []
    for c in COLUMN_ORDER:  # canonical order regardless of request order, like the PDF
        if c in columns:
            cols.extend(["upgrades_used", "upgrades_limit"] if c == "upgrades" else [c])
    labels = {**COLUMN_LABELS, "upgrades_used": "Upgrades Used", "upgrades_limit": "Upgrades Limit"}

    wb = Workbook()
    ws = wb.active
    ws.title = "Customers"
    ws.append([labels[c] for c in cols])
    for cell in ws[1]:
        cell.font, cell.fill = _HEADER_FONT, _HEADER_FILL
        cell.alignment = Alignment(vertical="center")

    for row in customers:
        values = [_cell_value(c, row) for c in cols]
        ws.append([v for v, _color, _fmt in values])
        for cell, (_v, color, fmt) in zip(ws[ws.max_row], values):
            if color:
                cell.font = Font(color=color)
            if fmt:
                cell.number_format = fmt

    for i, c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = _WIDTHS.get(c, 14)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions

    info = wb.create_sheet("Export info")
    for line in (
        ("Sedna Ops — Customer Version List",),
        ("Generated", generated_at),
        ("Customers", len(customers)),
        ("Note", "Renewal and Health are never included in this export."),
    ):
        info.append(line)
    info["A1"].font = Font(bold=True, size=13)
    info.column_dimensions["A"].width = 14
    info.column_dimensions["B"].width = 60

    wb.properties.title = "Sedna Ops — Customer Version List"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
